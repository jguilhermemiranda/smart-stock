from __future__ import annotations

import json
import sqlite3
from pathlib import Path
from typing import Any

from database.db import Database


class ExcelError(ValueError):
    """Raised when an Excel file cannot be read or validated."""


class ExcelService:
    SHEETS = ("Cards", "Drawers", "Inventory", "Permissions")

    def export(self, database: Database, path: str | Path) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError as exc:
            raise ExcelError("A integração com Excel requer o pacote openpyxl.") from exc

        workbook = Workbook()
        workbook.remove(workbook.active)
        header_fill = PatternFill("solid", fgColor="2C3871")
        header_font = Font(color="FFFFFF", bold=True)

        sheets = {
            "Cards": (
                ("uid", "owner_name", "status", "attributes"),
                self._card_rows(database),
            ),
            "Drawers": (
                ("label", "x_mm", "y_mm", "z_mm", "calibrated"),
                self._drawer_rows(database),
            ),
            "Inventory": (
                ("drawer_label", "item_name", "quantity"),
                self._inventory_rows(database),
            ),
            "Permissions": (
                ("card_uid", "drawer_label"),
                self._permission_rows(database),
            ),
        }
        for name, (headers, rows) in sheets.items():
            sheet = workbook.create_sheet(name)
            sheet.append(headers)
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
            for row in rows:
                sheet.append(row)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 42)
                sheet.column_dimensions[column[0].column_letter].width = width
        workbook.save(path)

    def import_file(self, database: Database, path: str | Path) -> dict[str, int]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ExcelError("A integração com Excel requer o pacote openpyxl.") from exc

        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            self._check_sheets(workbook.sheetnames)
            cards = self._read_rows(workbook["Cards"], ("uid", "owner_name", "status", "attributes"))
            drawers = self._read_rows(workbook["Drawers"], ("label", "x_mm", "y_mm", "z_mm", "calibrated"))
            inventory = self._read_rows(workbook["Inventory"], ("drawer_label", "item_name", "quantity"))
            permissions = self._read_rows(workbook["Permissions"], ("card_uid", "drawer_label"))
            return self._replace_database(database, cards, drawers, inventory, permissions)
        except ExcelError:
            raise
        except (OSError, KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
            raise ExcelError(f"Não foi possível ler o arquivo Excel: {exc}") from exc

    @staticmethod
    def _check_sheets(sheetnames: list[str]) -> None:
        missing = [name for name in ExcelService.SHEETS if name not in sheetnames]
        if missing:
            raise ExcelError(f"Abas obrigatórias ausentes: {', '.join(missing)}.")

    @staticmethod
    def _read_rows(sheet: Any, headers: tuple[str, ...]) -> list[dict[str, Any]]:
        rows = sheet.iter_rows(values_only=True)
        actual_headers = tuple(str(value or "").strip() for value in next(rows, ()))
        if actual_headers != headers:
            raise ExcelError(f"A aba {sheet.title} deve começar com: {', '.join(headers)}.")
        return [dict(zip(headers, row)) for row in rows if any(value is not None for value in row)]

    @staticmethod
    def _required(row: dict[str, Any], field: str, sheet: str) -> Any:
        value = row.get(field)
        if value is None or str(value).strip() == "":
            raise ExcelError(f"Campo obrigatório '{field}' vazio na aba {sheet}.")
        return value

    def _replace_database(
        self,
        database: Database,
        cards: list[dict[str, Any]],
        drawers: list[dict[str, Any]],
        inventory: list[dict[str, Any]],
        permissions: list[dict[str, Any]],
    ) -> dict[str, int]:
        connection = database.connection
        drawer_ids: dict[str, int] = {}
        card_ids: dict[str, int] = {}
        try:
            connection.execute("BEGIN")
            connection.execute("DELETE FROM card_drawer_permissions")
            connection.execute("DELETE FROM drawer_inventory")
            connection.execute("DELETE FROM cards")
            connection.execute("DELETE FROM drawers")
            for row in drawers:
                label = str(self._required(row, "label", "Drawers")).strip()
                calibrated = bool(row.get("calibrated", False))
                coordinates = tuple(float(self._required(row, axis, "Drawers")) for axis in ("x_mm", "y_mm", "z_mm")) if calibrated else (None, None, None)
                cursor = connection.execute(
                    "INSERT INTO drawers (label, pos_x_mm, pos_y_mm, pos_z_mm, calibrated) VALUES (?, ?, ?, ?, ?)",
                    (label, *coordinates, int(calibrated)),
                )
                drawer_ids[label] = cursor.lastrowid
            for row in cards:
                uid = str(self._required(row, "uid", "Cards")).replace(" ", "").upper()
                attributes = row.get("attributes") or "{}"
                if isinstance(attributes, dict):
                    attributes = json.dumps(attributes, ensure_ascii=False)
                json.loads(str(attributes))
                cursor = connection.execute(
                    "INSERT INTO cards (uid, owner_name, status, attributes_json) VALUES (?, ?, ?, ?)",
                    (uid, str(self._required(row, "owner_name", "Cards")).strip(), str(row.get("status") or "active"), str(attributes)),
                )
                card_ids[uid] = cursor.lastrowid
            for row in inventory:
                drawer_label = str(self._required(row, "drawer_label", "Inventory")).strip()
                if drawer_label not in drawer_ids:
                    raise ExcelError(f"Gaveta não encontrada na aba Inventory: {drawer_label}.")
                connection.execute(
                    "INSERT INTO drawer_inventory (drawer_id, item_name, quantity) VALUES (?, ?, ?)",
                    (drawer_ids[drawer_label], str(self._required(row, "item_name", "Inventory")).strip(), int(self._required(row, "quantity", "Inventory"))),
                )
            for row in permissions:
                card_uid = str(self._required(row, "card_uid", "Permissions")).replace(" ", "").upper()
                drawer_label = str(self._required(row, "drawer_label", "Permissions")).strip()
                if card_uid not in card_ids or drawer_label not in drawer_ids:
                    raise ExcelError("Permissão referencia cartão ou gaveta inexistente.")
                connection.execute(
                    "INSERT INTO card_drawer_permissions (card_id, drawer_id) VALUES (?, ?)",
                    (card_ids[card_uid], drawer_ids[drawer_label]),
                )
            connection.commit()
        except ExcelError:
            connection.rollback()
            raise
        except (sqlite3.IntegrityError, TypeError, ValueError) as exc:
            connection.rollback()
            raise ExcelError(f"Dados inválidos para importar: {exc}") from exc
        return {"cards": len(cards), "drawers": len(drawers), "inventory": len(inventory), "permissions": len(permissions)}

    @staticmethod
    def _card_rows(database: Database) -> list[tuple[Any, ...]]:
        return [(row["uid"], row["owner_name"], row["status"], row["attributes_json"]) for row in database.rows("SELECT uid, owner_name, status, attributes_json FROM cards ORDER BY id")]

    @staticmethod
    def _drawer_rows(database: Database) -> list[tuple[Any, ...]]:
        return [(row["label"], row["pos_x_mm"], row["pos_y_mm"], row["pos_z_mm"], bool(row["calibrated"])) for row in database.rows("SELECT label, pos_x_mm, pos_y_mm, pos_z_mm, calibrated FROM drawers ORDER BY id")]

    @staticmethod
    def _inventory_rows(database: Database) -> list[tuple[Any, ...]]:
        return [(row["label"], item["item_name"], item["quantity"]) for row in database.rows("SELECT id, label FROM drawers ORDER BY id") for item in database.drawer_inventory(row["id"])]

    @staticmethod
    def _permission_rows(database: Database) -> list[tuple[Any, ...]]:
        return [(card["uid"], drawer["label"]) for card in database.rows("SELECT uid, id FROM cards ORDER BY id") for drawer in database.rows("SELECT label FROM drawers JOIN card_drawer_permissions p ON p.drawer_id = drawers.id WHERE p.card_id = ? ORDER BY drawers.id", (card["id"],))]

